import argparse
import csv
import json
import os

import openpyxl

INSTRUCTOR_CODES = ["Lec","RtW","FUp","PQ","CQ","AnQ","MG","1o1","D/V","Adm","W","O"]
INSTRUCTOR_START_COL = 2  # Column B
DATA_START_ROW = 4        # Row 4 = first coded window
LABEL_COL = 1             # Column A, the "0-2" minute label

# A generated sheet stamps which lecture and which windows it belongs to, off to
# the right of row 1. Uploading three lectures at once pairs video to sheet by
# position; without this, swapping two sheets of similar-length lectures produces
# matching row counts AND matching minute labels, so nothing would catch it.
META_ROW = 1
META_LABEL_COL = 19       # Column S
META_VALUE_COL = 20       # Column T
META_MARKER = "COPUS_SHEET_META (do not edit)"


def read_template_meta(source):
    """
    The {lecture_id, windows, window_seconds} a generated sheet was built for.

    Returns None for a hand-made sheet with no stamp — those still work, they
    just can't be identity-checked.
    """
    if hasattr(source, "cell"):
        ws = source
    else:
        ws = openpyxl.load_workbook(source, data_only=True)["COPUS Coding Sheet"]
    if ws.cell(row=META_ROW, column=META_LABEL_COL).value != META_MARKER:
        return None
    raw = ws.cell(row=META_ROW, column=META_VALUE_COL).value
    try:
        return json.loads(raw)
    except (TypeError, ValueError):
        return None


def check_sheet_identity(xlsx_path, lecture_id=None, windows=None):
    """
    Compare a sheet's stamp against the lecture it is being used for.

    Returns a list of human-readable problems — empty when it matches or when
    the sheet carries no stamp.
    """
    meta = read_template_meta(xlsx_path)
    if meta is None:
        return []
    problems = []
    got_id = meta.get("lecture_id")
    if lecture_id and got_id and got_id != lecture_id:
        problems.append(
            f"sheet was generated for lecture '{got_id}' but is being used as "
            f"'{lecture_id}'")
    got_windows = meta.get("windows")
    if windows is not None and got_windows is not None:
        if sorted(got_windows) != sorted(windows):
            problems.append(
                f"sheet covers {len(got_windows)} windows "
                f"({got_windows[0]}..{got_windows[-1]}) but this lecture is being "
                f"classified on {len(windows)} "
                f"({sorted(windows)[0]}..{sorted(windows)[-1]})")
    return problems


def count_data_rows(ws):
    """
    How many coded window rows the sheet actually has.

    Replaces a hardcoded 30. A 30-row template silently truncated any lecture
    longer than 60 min (a 75-min lecture lost its last 15 min of ground truth)
    and padded any shorter one with blank rows that look like real "coded
    nothing" windows.
    """
    n = 0
    row = DATA_START_ROW
    while row <= ws.max_row:
        label = ws.cell(row=row, column=LABEL_COL).value
        has_code = any(
            ws.cell(row=row, column=INSTRUCTOR_START_COL + i).value
            for i in range(len(INSTRUCTOR_CODES))
        )
        if label is None and not has_code:
            break
        n += 1
        row += 1
    return n


def convert(excel_path, lecture_id, output_csv, windows=None, window_seconds=120):
    """
    Reads a filled COPUS instructor-only Excel template
    and converts it to long-format CSV for the validator.

    windows: the chunk indices these sheet rows correspond to, in order. Under
    sparse sampling the AI classifies non-contiguous chunks (e.g. 0-7, 17-24,
    35-42, 52-59), so sheet row DATA_START_ROW+i is window_index windows[i], NOT
    i. Getting this wrong makes kappa compare mismatched windows silently, so
    run.py always passes the same list it gave the classifier. When omitted, rows
    map to contiguous indices 0..n-1.
    """
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb["COPUS Coding Sheet"]

    for problem in check_sheet_identity(excel_path, lecture_id, windows):
        print(f"[WARN] {os.path.basename(excel_path)}: {problem}")

    n_rows = count_data_rows(ws)
    if windows is None:
        indices = list(range(n_rows))
    else:
        indices = list(windows)
        if len(indices) > n_rows:
            print(f"[warn] {os.path.basename(excel_path)}: asked for "
                  f"{len(indices)} windows but the sheet only has {n_rows} coded "
                  f"rows; using the first {n_rows}.")
            indices = indices[:n_rows]
        elif len(indices) < n_rows:
            print(f"[warn] {os.path.basename(excel_path)}: sheet has {n_rows} "
                  f"coded rows but only {len(indices)} windows were requested; "
                  f"ignoring the extra {n_rows - len(indices)} row(s).")

    rows = []
    mismatches = 0
    for i, window_idx in enumerate(indices):
        row = DATA_START_ROW + i
        codes_present = []

        for j, code in enumerate(INSTRUCTOR_CODES):
            col = INSTRUCTOR_START_COL + j
            val = ws.cell(row=row, column=col).value
            if val and str(val).strip().upper() == "X":
                codes_present.append(code)

        # Cross-check the sheet's own minute label against the chunk this row is
        # being mapped to -- catches a sheet pasted into the wrong rows.
        start_min = window_idx * window_seconds // 60
        end_min = (window_idx + 1) * window_seconds // 60
        label = ws.cell(row=row, column=LABEL_COL).value
        expected = f"{start_min}-{end_min}"
        if label is not None and str(label).strip() != expected:
            mismatches += 1
            print(f"[warn] row {row}: sheet says '{label}' but window "
                  f"{window_idx} is minutes {expected}")

        rows.append({
            "lecture_id": lecture_id,
            "window_index": window_idx,
            "window_start": window_idx * window_seconds,
            "window_end": (window_idx + 1) * window_seconds,
            "copus_codes": "|".join(codes_present)
        })

    os.makedirs(os.path.dirname(output_csv) if os.path.dirname(output_csv) else ".", exist_ok=True)

    with open(output_csv, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["lecture_id","window_index","window_start","window_end","copus_codes"])
        writer.writeheader()
        writer.writerows(rows)

    print(f"Converted {len(rows)} windows -> {output_csv}")
    if mismatches:
        print(f"[warn] {mismatches} row(s) disagree with the sheet's own minute "
              f"labels -- check the sheet is coded on the sampled windows.")
    print(f"  row {DATA_START_ROW}..{DATA_START_ROW + len(rows) - 1} -> "
          f"window_index {rows[0]['window_index']}..{rows[-1]['window_index']}"
          if rows else "  (no rows)")
    for r in rows:
        if r["copus_codes"]:
            print(f"  Window {r['window_index']:02d} "
                  f"({r['window_start'] // 60}-{r['window_end'] // 60} min): "
                  f"{r['copus_codes']}")

def main():
    parser = argparse.ArgumentParser(
        description="Convert a filled COPUS .xlsx to the long-format CSV the validator reads"
    )
    parser.add_argument("excel_path", help="Filled COPUS .xlsx")
    parser.add_argument("lecture_id", help="Lecture id, e.g. lecture_001")
    parser.add_argument("output_csv", help="Output .csv path")
    parser.add_argument("--windows", default=None,
                        help="Comma-separated chunk indices these sheet rows correspond "
                             "to. REQUIRED for a sparsely-sampled sheet -- without it the "
                             "rows are labelled 0..n-1 and kappa compares the wrong "
                             "windows. Use the same list you passed to --windows.")
    parser.add_argument("--window-seconds", type=int, default=120)
    args = parser.parse_args()

    windows = None
    if args.windows:
        try:
            windows = sorted({int(w) for w in args.windows.split(",") if w.strip()})
        except ValueError:
            parser.error(f"--windows must be comma-separated integers, got: {args.windows}")

    convert(args.excel_path, args.lecture_id, args.output_csv,
            windows=windows, window_seconds=args.window_seconds)


if __name__ == "__main__":
    main()
