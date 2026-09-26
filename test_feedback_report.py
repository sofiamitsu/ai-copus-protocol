"""
Acceptance tests for the Faculty Feedback Report rebuild (Change 9) and the
CUCEI score loader. No API calls: the Gemini narrative is stubbed out.

    uv run python test_feedback_report.py
"""
import os
import tempfile

import pandas as pd

import report.pdf_report as pr
from report.feedback_sections import (
    code_shares, cucei_profile, golden_profile, interpret_dimension,
    linking_observations, resolve_identity,
)
from utils.cucei import DIMENSIONS, load_all, load_scores
from utils.professor_ids import upsert_mapping

pr._gemini_text = lambda *a, **k: None  # never call Gemini from a test

tmp = tempfile.mkdtemp()


def write_cucei(path, rows):
    pd.DataFrame(rows).to_csv(path, index=False)
    return path


def scores(pid, means, n=20):
    return [{"professor_id": pid, "dimension": d, "mean": m, "sd": 0.5,
             "n_respondents": n} for d, m in zip(DIMENSIONS, means)]


# --- CUCEI loader: accepts a good file, rejects every malformed one -----------
good = write_cucei(os.path.join(tmp, "good.csv"),
                   scores("professor_1", [3.0] * 7) + scores("professor_2", [2.0] * 7))
assert load_all(os.path.join(tmp, "missing.csv")) is None
assert len(load_all(good)) == 14
assert load_scores("professor_404", good) is None
assert list(load_scores("professor_1", good)["dimension"]) == DIMENSIONS
print("ok  loader: missing file -> None, unknown professor -> None, dimension order kept")


for label, rows, needle in [
    ("workbook's broken Mean column", scores("professor_1", [24.57] * 7), "1-4 scale"),
    ("misspelled dimension", [{**scores("professor_1", [3.0])[0], "dimension": "Personalisation"}],
     "unknown CUCEI dimension"),
    ("duplicate row", scores("professor_1", [3.0] * 7) + scores("professor_1", [3.0]),
     "duplicate"),
]:
    bad = write_cucei(os.path.join(tmp, "bad.csv"), rows)
    try:
        load_all(bad)
        raise AssertionError(f"accepted a file with a {label}")
    except ValueError as e:
        assert needle in str(e), e
print("ok  loader rejects out-of-scale means (e.g. 24.57), misspelled dimensions, duplicates")

bad = os.path.join(tmp, "bad_cols.csv")
pd.DataFrame([{"professor_id": "p", "dimension": "Innovation", "mean": 3}]).to_csv(bad, index=False)
try:
    load_all(bad)
    raise AssertionError("accepted a file with no sd / n_respondents")
except ValueError as e:
    assert "missing column" in str(e)
print("ok  loader rejects a file missing required columns")

# --- Interpretations compare only to the scale midpoint --------------------------
assert "above the scale midpoint (2.5)" in interpret_dimension("Innovation", 3.0)
assert "below the scale midpoint (2.5)" in interpret_dimension("Innovation", 2.0)
assert "in line with the scale midpoint" in interpret_dimension("Innovation", 2.55)
print("ok  interpretation: vs the scale midpoint only")

# --- Linking observations ----------------------------------------------------
data_dir = os.path.join(tmp, "data")
os.makedirs(data_dir)
write_cucei(os.path.join(data_dir, "cucei_scores.csv"),
            scores("professor_1", [3.8, 2.2, 3.0, 3.0, 3.0, 3.6, 2.4])
            + scores("professor_2", [1.0] * 7))
rows = cucei_profile("professor_1", data_dir)
assert len(rows) == 7 and rows[0]["n"] == 20
# Another professor in the same data must not change a word of this report.
alone_dir = os.path.join(tmp, "data_alone")
os.makedirs(alone_dir)
write_cucei(os.path.join(alone_dir, "cucei_scores.csv"),
            scores("professor_1", [3.8, 2.2, 3.0, 3.0, 3.0, 3.6, 2.4]))
assert cucei_profile("professor_1", alone_dir) == rows
assert not any("average" in r["interpretation"] for r in rows)
assert not any(k in rows[0] for k in ("study_avg", "n_courses"))
print("ok  CUCEI section uses only this professor's scores; others change nothing")

shares = {"Lec": 95.0, "MG": 18.0, "AnQ": 40.0, "PQ": 30.0, "D/V": 10.0}
obs = linking_observations(shares, rows)
assert len(obs) == 3, obs
assert obs[0].startswith("You answered student questions (AnQ) in 40%"), obs[0]
assert not any("(PQ)" in o for o in obs), "one observation per CUCEI dimension"
assert not any("(Lec)" in o for o in obs), "lecturing is a fallback only"
assert "It may be worth asking students" in obs[0]           # Involvement 2.2 is below
assert "may be part of why students rate Personalization favorably" in obs[1]
assert all(" because " not in o and "caused" not in o for o in obs)
print("ok  linking: most frequent first, one per dimension, Lec skipped, hedged wording")

only_lec = linking_observations({"Lec": 100.0}, rows)
assert len(only_lec) == 1 and "(Lec)" in only_lec[0]
assert linking_observations(shares, None) == []
print("ok  linking: Lec used only when nothing else was seen; no CUCEI -> no observations")

# --- Golden aggregate -----------------------------------------------------------
assert golden_profile(os.path.join(tmp, "nothing_here")) is None
for name, codes in (("g1/lec_a", ["Lec", "Lec|MG", "Lec|PQ", "Lec"]),
                    ("g2/lec_b", ["Lec|MG", "MG", "Lec", "Lec|AnQ"])):
    d = os.path.join(data_dir, "golden", name)
    os.makedirs(d)
    pd.DataFrame({"copus_codes": codes, "window_start": 0, "window_end": 120}).to_csv(
        os.path.join(d, "results_multimodal.csv"), index=False)
g = golden_profile(data_dir)
assert g["n_lectures"] == 2 and g["n_windows"] == 8
assert g["shares"]["MG"] == 37.5 and g["shares"]["Lec"] == 87.5, g["shares"]
assert code_shares(pd.DataFrame({"copus_codes": ["Lec|Lec", None]})) == {"Lec": 50.0}
print("ok  golden: pools every results_multimodal.csv under data/golden, window-weighted")

# --- Identity: sidebar wins, mapping fills blanks -------------------------------
out = os.path.join(tmp, "out")
upsert_mapping(out, "lec1", "professor_1", "Dr. Rodrigo", "EGN 3000")
assert resolve_identity(out, {}) == ("professor_1", "Dr. Rodrigo", "EGN 3000")
assert resolve_identity(out, {"professor": "Dr. R", "course": ""}) == (
    "professor_1", "Dr. R", "EGN 3000")
print("ok  identity: mapping fills blank sidebar fields -> no more 'Professor ( , )'")

# --- Full PDF build, with and without CUCEI / golden -----------------------------
lec_dir = os.path.join(out, "lec1")
os.makedirs(lec_dir)
pd.DataFrame({
    "lecture_id": "lec1", "professor_id": "professor_1",
    "window_index": range(6), "window_start": [i * 120 for i in range(6)],
    "window_end": [(i + 1) * 120 for i in range(6)], "arm": "multimodal",
    "model": "m", "copus_codes": ["Lec", "Lec|AnQ", "Lec|MG", "Lec|PQ", "Lec", "Adm"],
    "reasoning": "",
}).to_csv(os.path.join(lec_dir, "results_multimodal.csv"), index=False)
kappa = os.path.join(out, "combined_kappa.csv")
pd.DataFrame([
    {"professor_id": "professor_1", "code": "Lec", "n_total_windows": 6, "n_human_marked": 5,
     "n_ai_marked": 5, "pct_agreement": 100.0, "ai_vs_sofia_kappa": 1.0, "ai_vs_sofia_ac1": 1.0},
    {"professor_id": "professor_1", "code": "O", "n_total_windows": 6, "n_human_marked": 0,
     "n_ai_marked": 1, "pct_agreement": 83.3, "ai_vs_sofia_kappa": 0.0, "ai_vs_sofia_ac1": 0.8},
    {"professor_id": "professor_1", "code": "codes_clearing_kappa_0.7",
     "n_codes_human_observed": 1, "ai_vs_sofia_kappa": 1},
    {"professor_id": "professor_1", "code": "codes_clearing_ac1_0.7",
     "n_codes_human_observed": 1, "ai_vs_sofia_ac1": 1},
]).to_csv(kappa, index=False)

for label, dd in (("full", data_dir), ("placeholders", os.path.join(tmp, "empty"))):
    pdf = os.path.join(tmp, f"{label}.pdf")
    pr.generate_faculty_report(
        professor_name="Dr. Rodrigo", course_name="EGN 3000", semester="",
        lecture_results=pr._discover_lectures(out), kappa_results=kappa,
        survey_data=None, output_path=pdf, professor_id="professor_1", data_dir=dd)
    assert os.path.getsize(pdf) > 5000, label
    with open(pdf, "rb") as f:
        assert f.read(5) == b"%PDF-"
print("ok  report builds with CUCEI + golden, and with neither (placeholders)")

clearing = pr._clearing_summary(pd.read_csv(kappa))
assert "1 of 1" in clearing, clearing
print("ok  validation summary reads the threshold counts over human-observed codes")

# --- Recommendations use CUCEI, in both the fallback and the Gemini prompt -------
from report.feedback_sections import cucei_recommendation

freq = pr.code_frequency(pd.DataFrame({"copus_codes": ["Lec", "Lec|AnQ", "Lec"]}))
lo_shares = {"Lec": 100.0, "AnQ": 12.5}

# professor_1 (from above): Involvement 2.2 is furthest below the midpoint.
rec = cucei_recommendation(lo_shares, rows)
assert rec.startswith("Students rated Involvement 2.20/4, below the scale midpoint"), rec
assert "AnQ in 12.5%" in rec and "PQ in 0%" in rec and "Lec" not in rec, rec
assert "not a diagnosis" in rec
print("ok  CUCEI rec: dimension furthest below the midpoint, with this professor's own rates")

fallback = pr.recommendations(freq, 20.0, shares=lo_shares, cucei_rows=rows)
assert fallback[0] == rec and len(fallback) <= 3, fallback
assert pr.recommendations(freq, 20.0) == pr._fallback_recommendations(freq, 20.0)
print("ok  fallback leads with the CUCEI suggestion; unchanged when there is no CUCEI")

# Linked behavior already common -> must NOT recommend doing more of it.
common = cucei_recommendation({"AnQ": 60.0}, rows)
assert "already common in your lectures (AnQ in 60%" in common, common
assert "unlikely to be the answer" in common and "gives more students a way" not in common
print("ok  linked behavior already common (>=25%) -> no 'do more of it' advice")

all_high = [{**r, "mean": 3.9} for r in rows]
assert cucei_recommendation(lo_shares, all_high) is None
print("ok  no dimension below its comparison point -> no CUCEI-based suggestion")

# Dimensions with no linked instructor behavior never produce a suggestion.
only_sat = [{**r, "mean": (1.2 if r["dimension"] == "Satisfaction" else 3.9)} for r in rows]
assert cucei_recommendation(lo_shares, only_sat) is None
print("ok  Satisfaction low but no linked COPUS behavior -> no suggestion invented")

captured = {}
pr._gemini_text = lambda prompt, **k: captured.setdefault("prompt", prompt) and None
pr.recommendations(freq, 20.0, shares=lo_shares, cucei_rows=rows)
p = captured["prompt"]
assert "CUCEI Involvement: 2.20/4 (n=20), below the scale midpoint" in p, p
assert "average" not in p.split("Student perceptions")[1], p
assert "AT MOST ONE" in p and "do not claim that any behavior caused it" in p
pr._gemini_text = lambda *a, **k: None
print("ok  Gemini prompt gets each score vs its comparison, capped at one CUCEI-tied rec")

# --- Reading filled CUCEI workbooks straight from data/cucei/ --------------------
import openpyxl
from utils.cucei import (
    load_workbook_dir, load_workbook_scores, professor_id_from_filename,
)

assert professor_id_from_filename("CUCEI PROFESSOR 1 (1).xlsm") == "professor_1"
assert professor_id_from_filename("cucei_professor_12.xlsx") == "professor_12"
assert professor_id_from_filename("CUCEI final.xlsm") is None
print("ok  professor_id parsed from the workbook filename")

ANSWERS = {  # 2 respondents; every item answered
    1: ["SA", "A"], 2: ["D", "SD"],
}


def fake_workbook(path, means=None, answers=("SA", "A")):
    """
    A minimal stand-in for the scoring tool: Scoring Key, Data Entry (raw answers,
    no formulas) and Group Summary (plain numbers, as Excel would have cached).
    `means` overrides Group Summary so a disagreement can be simulated.
    """
    wb = openpyxl.Workbook()
    key = wb.active
    key.title = "Scoring Key"
    key.append(["Subscale", "Positive", "Negative"])
    items = iter(range(1, 50))
    layout = {}
    for dim in DIMENSIONS:
        pos = [next(items) for _ in range(4)]
        neg = [next(items) for _ in range(3)]
        layout[dim] = (pos, neg)
        key.append([dim, ", ".join(map(str, pos)), ", ".join(map(str, neg))])

    de = wb.create_sheet("Data Entry")
    de.append(["Respondent ID"] + [f"Q{i}" for i in range(1, 50)])
    for r, a in enumerate(answers, start=1):
        de.append([f"S{r}"] + [a] * 49)

    # Truth: positives score SA=4/A=3, negatives reverse -> (4*p + 3*n) / 7
    pos_score = {"SA": 4, "A": 3, "D": 2, "SD": 1}
    truth = [sum([pos_score[a]] * 4 + [5 - pos_score[a]] * 3) / 7 for a in answers]
    true_mean = sum(truth) / len(truth)

    gs = wb.create_sheet("Group Summary")
    gs.append(["CUCEI Group Summary"])
    gs.append(["Subscale", "Complete N", "Mean", "Std. Dev."])
    for dim in DIMENSIONS:
        gs.append([dim, len(answers), (means or {}).get(dim, true_mean), 0.1])
    wb.save(path)
    return path, true_mean


wb_dir = os.path.join(tmp, "cucei_books")
os.makedirs(wb_dir)
good_wb, true_mean = fake_workbook(os.path.join(wb_dir, "CUCEI PROFESSOR 7.xlsm"))
df_wb = load_workbook_scores(good_wb)
assert list(df_wb["professor_id"].unique()) == ["professor_7"]
assert list(df_wb["dimension"]) == DIMENSIONS
assert abs(df_wb["mean"].iloc[0] - true_mean) < 1e-9
assert df_wb["n_respondents"].iloc[0] == 2
print(f"ok  workbook read: 7 dimensions, mean {true_mean:.3f} matches the raw answers")

# Group Summary that disagrees with the sheet's own answers must be refused.
bad_wb, _ = fake_workbook(os.path.join(tmp, "CUCEI PROFESSOR 8.xlsm"),
                          means={"Satisfaction": 17.58})
try:
    load_workbook_scores(bad_wb)
    raise AssertionError("accepted a workbook whose Group Summary contradicts its data")
except ValueError as e:
    assert "Satisfaction" in str(e) and "Sum / 7" in str(e), e
print("ok  workbook whose totals disagree with its answers is refused (the 17.58 bug)")

assert load_workbook_dir(os.path.join(tmp, "no_such_dir")) is None
assert len(load_workbook_dir(wb_dir)) == 7
print("ok  folder scan: empty folder -> None, one workbook -> 7 rows")

# Workbooks and a CSV are read together; the same professor in both is an error.
both = os.path.join(tmp, "both")
os.makedirs(both)
write_cucei(os.path.join(both, "cucei_scores.csv"), scores("professor_3", [3.0] * 7))
merged = load_all(os.path.join(both, "cucei_scores.csv"), wb_dir)
assert sorted(merged["professor_id"].unique()) == ["professor_3", "professor_7"]
write_cucei(os.path.join(both, "clash.csv"), scores("professor_7", [3.0] * 7))
try:
    load_all(os.path.join(both, "clash.csv"), wb_dir)
    raise AssertionError("accepted professor_7 in both a workbook and a CSV")
except ValueError as e:
    assert "duplicate" in str(e), e
print("ok  workbooks + CSV merge; a professor in both sources is rejected")

# --- Streamlit uploads: staged alone, never mixed with other professors ----------
from utils.cucei import stage_data_dir

base = os.path.join(tmp, "base_data")
os.makedirs(os.path.join(base, "cucei"))
fake_workbook(os.path.join(base, "cucei", "CUCEI PROFESSOR 9.xlsm"))
write_cucei(os.path.join(base, "cucei_scores.csv"), scores("professor_3", [3.0] * 7))
os.makedirs(os.path.join(base, "golden", "lec1"))
upload_dir = os.path.join(tmp, "uploads")
os.makedirs(upload_dir)
# The file name need not carry the professor number when the caller supplies it.
upload, upload_mean = fake_workbook(os.path.join(upload_dir, "CUCEI final.xlsm"))
uploaded = load_workbook_scores(upload, "professor_7")

staged = stage_data_dir(uploaded, os.path.join(tmp, "staged"), base)
merged = load_all(os.path.join(staged, "cucei_scores.csv"), os.path.join(staged, "cucei"))
assert list(merged["professor_id"].unique()) == ["professor_7"], merged
assert abs(merged["mean"].iloc[0] - upload_mean) < 1e-9
assert os.path.isdir(os.path.join(staged, "golden", "lec1"))
assert cucei_profile("professor_7", staged)[0]["mean"] == merged["mean"].iloc[0]
assert os.path.isdir(os.path.join(stage_data_dir(uploaded, os.path.join(tmp, "s2"), None)))
print("ok  staged upload holds only that professor's scores (+ golden), none from data/")

print("\nAll feedback report tests passed.")
